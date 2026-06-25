"""
Hard test cases for document_parser.py: corrupted files, encrypted PDFs,
oversized files, empty files, image-only PDFs, and malformed Excel/Word
files. Uses a minimal fake UploadedFile since we don't have Streamlit's
runtime here.
"""

import io
import sys

sys.path.insert(0, "/home/claude/build")
import document_parser  # noqa: E402

import fitz
import docx as docx_lib
import openpyxl

passed = 0
failed = 0


def check(label, condition, detail=""):
    global passed, failed
    if condition:
        passed += 1
        print(f"  PASS: {label}")
    else:
        failed += 1
        print(f"  FAIL: {label}  {detail}")


class FakeUploadedFile:
    """Minimal stand-in for streamlit.runtime.uploaded_file_manager.UploadedFile."""
    def __init__(self, name, data: bytes):
        self.name = name
        self._buf = io.BytesIO(data)

    def read(self):
        return self._buf.read()

    def seek(self, pos):
        self._buf.seek(pos)


def make_real_pdf(text="Hello vendor risk world", encrypted=False, pages=1, no_text=False):
    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page()
        if not no_text:
            page.insert_text((50, 72), text)
    buf = io.BytesIO()
    if encrypted:
        doc.save(buf, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="ownerpw", user_pw="userpw")
    else:
        doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_real_docx(paragraphs=("Hello", "World")):
    doc = docx_lib.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_real_xlsx(data=(("Vendor", "Risk"), ("Acme", "High"))):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in data:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


print("=== PDF parsing ===")

text, err = document_parser.parse_file(FakeUploadedFile("good.pdf", make_real_pdf("Vendor ACME has weak MFA controls")))
check("valid pdf with text parses with no error", err is None and "ACME" in text, (text, err))

text, err = document_parser.parse_file(FakeUploadedFile("corrupt.pdf", b"%PDF-1.4 not a real pdf body !!!@#$"))
check("corrupted pdf returns clear error", err is not None and "pdf" in err.lower(), err)

text, err = document_parser.parse_file(FakeUploadedFile("empty.pdf", b""))
check("zero-byte file returns clear error", err is not None and "empty" in err.lower(), err)

encrypted_bytes = make_real_pdf("secret content", encrypted=True)
text, err = document_parser.parse_file(FakeUploadedFile("locked.pdf", encrypted_bytes))
check("password-protected pdf returns clear error (not silently empty)", err is not None and "password" in err.lower(), err)

image_only_bytes = make_real_pdf(no_text=True)
text, err = document_parser.parse_file(FakeUploadedFile("scanned.pdf", image_only_bytes))
check("image-only/scanned pdf flags no extractable text", err is not None and "no extractable text" in err.lower(), err)

oversized = b"x" * (document_parser.MAX_FILE_SIZE_BYTES + 1)
text, err = document_parser.parse_file(FakeUploadedFile("huge.pdf", oversized))
check("oversized file rejected before parsing", err is not None and "exceeds" in err.lower(), err)

many_pages_pdf = make_real_pdf("page text", pages=3)
text, err = document_parser.parse_file(FakeUploadedFile("multi.pdf", many_pages_pdf))
check("multi-page pdf parses all pages", err is None and text.count("page text") == 3, (text.count("page text"), err))


print("\n=== DOCX parsing ===")

text, err = document_parser.parse_file(FakeUploadedFile("good.docx", make_real_docx(["Vendor: Acme", "Risk: High"])))
check("valid docx parses with no error", err is None and "Acme" in text, (text, err))

text, err = document_parser.parse_file(FakeUploadedFile("corrupt.docx", b"PK\x03\x04 not really a zip/docx"))
check("corrupted docx returns clear error", err is not None and "word" in err.lower(), err)

text, err = document_parser.parse_file(FakeUploadedFile("blank.docx", make_real_docx([])))
check("docx with no paragraphs flags no readable text", err is not None and "no readable text" in err.lower(), err)


print("\n=== XLSX parsing ===")

text, err = document_parser.parse_file(FakeUploadedFile("good.xlsx", make_real_xlsx()))
check("valid xlsx parses with no error", err is None and "Acme" in text, (text, err))

text, err = document_parser.parse_file(FakeUploadedFile("corrupt.xlsx", b"not a real xlsx file at all"))
check("corrupted xlsx returns clear error", err is not None and "excel" in err.lower(), err)

empty_wb_bytes = io.BytesIO()
import openpyxl as _o
_wb = _o.Workbook()
_wb.active.title = "EmptySheet"
_wb.save(empty_wb_bytes)
text, err = document_parser.parse_file(FakeUploadedFile("empty.xlsx", empty_wb_bytes.getvalue()))
check("xlsx with empty sheet flags no data", err is not None and "no data" in err.lower(), err)


print("\n=== Unsupported / edge cases ===")

text, err = document_parser.parse_file(FakeUploadedFile("malware.exe", b"MZ\x90\x00fakebinarycontent"))
check("unsupported extension rejected", err is not None and "unsupported" in err.lower(), err)

text, err = document_parser.parse_file(FakeUploadedFile("noextension", b"some content"))
check("file with no extension rejected", err is not None, err)

text, err = document_parser.parse_file(FakeUploadedFile("file.PDF", make_real_pdf("uppercase extension test")))
check("uppercase extension still matched (case-insensitive)", err is None, err)

print(f"\n{'='*50}\nRESULTS: {passed} passed, {failed} failed\n{'='*50}")
sys.exit(1 if failed else 0)
